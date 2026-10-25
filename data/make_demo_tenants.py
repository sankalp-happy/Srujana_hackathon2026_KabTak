"""Demo tenants for "Kab Tak?" multi-tenancy (all data is made up; no real company names).

Writes, per docs/planning/multi-tenant-contract.md §1:
  data/tenants/shree-textiles/  power-loom fabric trader, Ichalkaranji, Marathi-speaking owner: no overdraft,
                                wedding-season peak (Nov-Feb), one slow and one cash-like buyer
  data/tenants/sai-snacks/      3-month-old snack maker: ~8 paid invoices, so the engine falls back to pooled/default curves
and a messy TallyPrime-style upload for the onboarding demo:
  data/samples/kolhapur_auto_spares_sales_register.{xlsx,csv} + README.md (which rows are deliberately bad)

Run `python3 data/make_demo_tenants.py` to rebuild everything and run the self-check.
"""
import csv
import hashlib
import json
import sys
from datetime import date, timedelta
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
from engine.ledger import GST_RATE, KINDS, TDS_RATE   # read-only constants

TODAY = date(2026, 10, 25)
SEED = 11
TENANTS = ROOT / "data" / "tenants"
SAMPLES = ROOT / "data" / "samples"
INV_COLS = ["invoice_no", "buyer", "invoice_date", "taxable_inr", "gst_inr", "gross_inr", "tds_inr",
            "paid_date", "disputed", "treds_accepted"]
OBL_COLS = ["date", "label", "kind", "amount_inr"]
POS_KEYS = {"bank_inr", "od_limit_inr", "od_drawn_inr", "stock_inr", "creditors_inr"}
VALID_KINDS = KINDS + ["other"]   # contract §1: KINDS gains "other"

# ---- Shree Textiles ----
SHREE_END = date(2027, 1, 31)       # obligations end here (as Mahesh's); invoices run a month further to feed them
SHREE_BUYERS = ["Patil Fabrics", "Mehta Wholesale", "Jain Garments", "Kothari Sarees", "Counter sales"]
SHREE_MEDIAN = {"Patil Fabrics": 45, "Mehta Wholesale": 75, "Jain Garments": 50, "Kothari Sarees": 30, "Counter sales": 1}
SHREE_SIGMA = {"Patil Fabrics": 0.25, "Mehta Wholesale": 0.25, "Jain Garments": 0.30, "Kothari Sarees": 0.25}
SHREE_SHARE = {"Patil Fabrics": 0.28, "Mehta Wholesale": 0.22, "Jain Garments": 0.20, "Kothari Sarees": 0.12, "Counter sales": 0.18}
SHREE_N = {"Patil Fabrics": 3, "Mehta Wholesale": 2, "Jain Garments": 3, "Kothari Sarees": 2, "Counter sales": 8}   # invoices/month
SHREE_REVENUE = 1_300_000           # taxable Rs per month before seasonality
SEASON = {1: 1.4, 2: 1.2, 3: 0.9, 4: 0.85, 5: 0.85, 6: 0.8, 7: 0.7, 8: 0.75, 9: 0.85, 10: 1.0, 11: 1.5, 12: 1.5}
YARN_SHARE = 0.55                   # yarn as a share of next month's sales, bought ahead of the season
SHREE_PARAMS = {"gst_rate": 0.05, "tds_rate": 0.0}   # fabric is 5% GST; traders' buyers deduct no TDS
SAI_PARAMS = {"gst_rate": 0.12, "tds_rate": 0.0}     # namkeen/snacks at 12%


def _months(start, end):
    d = start
    while d <= end:
        yield d
        d = (d.replace(day=28) + timedelta(days=4)).replace(day=1)


def _month_end(m):
    return (m.replace(day=28) + timedelta(days=4)).replace(day=1) - timedelta(days=1)


def _finish(df, prefix, gst, tds):
    """Sort by date, then add GST/gross/TDS and invoice numbers (whole rupees, ledger-style rounding)."""
    df = df.sort_values("invoice_date", kind="stable").reset_index(drop=True)
    df["gst_inr"] = (df.taxable_inr * gst).round().astype(int)
    df["gross_inr"] = df.taxable_inr + df.gst_inr
    df["tds_inr"] = (df.taxable_inr * tds).round().astype(int)
    df["invoice_no"] = [f"{prefix}-{i + 1:04d}" for i in range(len(df))]
    return df[INV_COLS]


def shree_invoices(rng):
    rows = []
    for m in _months(date(2025, 5, 1), date(2027, 2, 1)):
        month_factor = SEASON[m.month] * rng.uniform(0.92, 1.08)
        last = _month_end(m).day
        for b in SHREE_BUYERS:
            w = rng.lognormal(0, 0.35, SHREE_N[b])
            for amt in SHREE_SHARE[b] * SHREE_REVENUE * month_factor * w / w.sum():
                d = m.replace(day=int(rng.integers(1, last + 1)))
                delay = int(rng.integers(0, 3)) if b == "Counter sales" else int(round(SHREE_MEDIAN[b] * rng.lognormal(0, SHREE_SIGMA[b])))
                rows.append([b, d, int(round(amt, -2)), delay])
    df = pd.DataFrame(rows, columns=["buyer", "invoice_date", "taxable_inr", "delay"])
    df["paid_date"] = [d + timedelta(days=int(x)) if d + timedelta(days=int(x)) <= TODAY else None
                       for d, x in zip(df.invoice_date, df.delay)]
    df["disputed"], df["treds_accepted"] = 0, 0
    # planted exception: an old Mehta invoice, disputed over a short-shipped lot, still open
    df = pd.concat([df, pd.DataFrame([["Mehta Wholesale", date(2026, 6, 18), 180_000, 0, None, 1, 0]], columns=df.columns)], ignore_index=True)
    return _finish(df, "SHT", SHREE_PARAMS["gst_rate"], SHREE_PARAMS["tds_rate"])


def shree_obligations(inv):
    """Known cash-outs TODAY+1 .. SHREE_END: weekly wages, yarn supplier, power, loom EMI, rent, GST on the 20th."""
    rows = []
    add = lambda d, label, kind, amt: rows.append([d, label, kind, int(round(amt, -3))]) if TODAY < d <= SHREE_END else None
    rev = inv.groupby(pd.to_datetime(inv.invoice_date).dt.to_period("M")).taxable_inr.sum()
    gst = SHREE_PARAMS["gst_rate"]
    for m in _months(date(2026, 9, 1), SHREE_END):
        season = SEASON[m.month] >= 1.2
        nxt = rev.get(pd.Period(m + timedelta(days=31), "M"), 0)
        add(date(m.year + (m.month == 12), m.month % 12 + 1, 20), "GST", "gst", gst * rev.get(pd.Period(m, "M"), 0) * (1 - YARN_SHARE))
        for day in (1, 15):   # yarn bought in two tranches a month, paid on 30-day terms
            add(m.replace(day=day) + timedelta(days=30), "Yarn (supplier)", "supplier", YARN_SHARE * nxt / 2 * (1 + gst))
        add(m.replace(day=12), "Power bill (looms)", "power", 165_000 if season else 140_000)
        add(m.replace(day=5), "Loom EMI", "emi", 65_000)
        add(m.replace(day=1), "Rent", "rent", 35_000)
    d = TODAY
    while d < SHREE_END:   # weavers are paid every Saturday
        d += timedelta(days=1)
        if d.weekday() == 5:
            add(d, "Weaver wages (weekly)", "wages", 65_000 if SEASON[d.month] >= 1.2 else 55_000)
    add(date(2026, 11, 5), "Diwali bonus (weavers)", "bonus", 120_000)
    return pd.DataFrame(sorted(rows), columns=OBL_COLS)


def _position(obl, bank, stock):
    # creditors = supplier bills falling due within 30 days, as in Mahesh's ledger
    cred = int(obl[(obl.kind == "supplier") & (obl.date <= TODAY + timedelta(days=30))].amount_inr.sum())
    return {"bank_inr": bank, "od_limit_inr": 0, "od_drawn_inr": 0, "stock_inr": stock, "creditors_inr": cred}


def shree():
    inv = shree_invoices(np.random.default_rng(SEED))
    obl = shree_obligations(inv)
    prof = {"tenant_id": "shree-textiles", "business_name": "Shree Textiles, Ichalkaranji", "buyers": SHREE_BUYERS,
            "today": TODAY.isoformat(), "synthetic": True, "params": SHREE_PARAMS}
    return prof, inv, obl, _position(obl, 1_100_000, 450_000)   # bank calibrated on the engine: worst week ~23% (mid-Dec)


# ---- Sai Snacks ----
SAI_BUYERS = ["Gupta Kirana", "Sweet Corner Bakery", "Daily Needs Mart"]
SAI_ROWS = [  # buyer, invoice date, taxable Rs, typical days to pay (None = still open / planned)
    ("Gupta Kirana", date(2026, 7, 27), 18_000, 18), ("Gupta Kirana", date(2026, 8, 24), 22_000, 22),
    ("Gupta Kirana", date(2026, 9, 21), 24_000, 19), ("Gupta Kirana", date(2026, 10, 12), 26_000, None),
    ("Sweet Corner Bakery", date(2026, 8, 3), 31_000, 38), ("Sweet Corner Bakery", date(2026, 9, 7), 28_000, 33),
    ("Sweet Corner Bakery", date(2026, 10, 5), 33_000, None),
    ("Daily Needs Mart", date(2026, 8, 10), 15_000, 25), ("Daily Needs Mart", date(2026, 8, 31), 17_500, 31),
    ("Daily Needs Mart", date(2026, 9, 19), 16_000, 27), ("Daily Needs Mart", date(2026, 10, 10), 19_000, None),
    ("Daily Needs Mart", date(2026, 10, 22), 21_000, None),
    ("Gupta Kirana", date(2026, 11, 2), 25_000, None), ("Sweet Corner Bakery", date(2026, 11, 9), 34_000, None),  # planned (Diwali orders)
]


def sai():
    rng = np.random.default_rng(SEED + 1)
    rows = []
    for b, d, amt, delay in SAI_ROWS:
        paid = d + timedelta(days=delay + int(rng.integers(-3, 4))) if delay else None
        rows.append([b, d, amt, paid, 0, 0])
    df = pd.DataFrame(rows, columns=["buyer", "invoice_date", "taxable_inr", "paid_date", "disputed", "treds_accepted"])
    assert all(p <= TODAY for p in df.paid_date.dropna())
    inv = _finish(df, "SAI", SAI_PARAMS["gst_rate"], SAI_PARAMS["tds_rate"])
    rows = []
    for m in _months(date(2026, 10, 1), date(2027, 1, 1)):
        for day, label, kind, amt in [(1, "Shop rent", "rent", 18_000), (5, "Helpers' wages", "wages", 28_000),
                                      (7, "Oven loan EMI", "emi", 12_000), (10, "Power", "power", 9_000),
                                      (12, "Gas cylinders", "other", 6_000), (15, "Flour & oil (supplier)", "supplier", 40_000)]:
            if TODAY < m.replace(day=day) <= date(2027, 1, 31):
                rows.append([m.replace(day=day), label, kind, amt])
    obl = pd.DataFrame(sorted(rows), columns=OBL_COLS)
    prof = {"tenant_id": "sai-snacks", "business_name": "Sai Snacks, Pune", "buyers": SAI_BUYERS,
            "today": TODAY.isoformat(), "synthetic": True, "params": SAI_PARAMS}
    return prof, inv, obl, _position(obl, 200_000, 35_000)   # bank calibrated on the engine: worst week ~17% (Jan)


def write_tenant(prof, inv, obl, pos):
    d = TENANTS / prof["tenant_id"]
    d.mkdir(parents=True, exist_ok=True)
    (d / "profile.json").write_text(json.dumps(prof, indent=2) + "\n")
    inv.to_csv(d / "invoices.csv", index=False)
    obl.to_csv(d / "obligations.csv", index=False)
    (d / "position.json").write_text(json.dumps(pos, indent=2) + "\n")


# ---- Kolhapur Auto Spares: messy TallyPrime sales register ----
TALLY_HEAD = ["Date", "Particulars", "Vch Type", "Vch No.", "Gross Total", "Received On"]
TALLY_BUYERS = {"Shinde Motors": (0.28, 30), "Kolhapur Garage Co": (0.22, 55), "Deshmukh Auto": (0.20, 42),
                "Patel Tyres & Spares": (0.18, 20), "Walk-in customers": (0.12, 2)}   # (share of rows, median days to pay)
TALLY_GST = 0.18


def inr(x, dp=2):
    """Indian digit grouping: 125000 -> '1,25,000.00'."""
    s = f"{abs(x):.{dp}f}"
    whole, frac = s.split(".")
    head, tail = whole[:-3], whole[-3:]
    while len(head) > 2:
        head, tail = head[:-2], head[-2:] + "," + tail
    return ("-" if x < 0 else "") + (head + "," if head else "") + tail + "." + frac


def tally_rows():
    """Returns (rows, notes): rows are string lists in TALLY_HEAD order; notes lists the deliberately bad rows."""
    rng = np.random.default_rng(SEED + 2)
    names, shares = list(TALLY_BUYERS), [v[0] for v in TALLY_BUYERS.values()]
    clean = []
    for m in _months(date(2026, 4, 1), date(2026, 10, 1)):
        n = int(rng.integers(12, 16))
        days = sorted(int(x) for x in rng.integers(1, _month_end(m).day + 1, n))
        for day in days:
            d = m.replace(day=day)
            if d > TODAY:
                continue
            b = str(rng.choice(names, p=shares))
            taxable = int(round(float(rng.lognormal(10.4, 0.8)), -2))   # median ~Rs 33k, long tail past Rs 1 L
            paid = d + timedelta(days=int(round(TALLY_BUYERS[b][1] * rng.lognormal(0, 0.3))))
            clean.append([d, b, round(taxable * (1 + TALLY_GST)), paid if paid <= TODAY else None])
    rows = [[d.strftime("%d-%m-%Y"), b, "Sales", f"SL/26-27/{i + 1:04d}", inr(g), p.strftime("%d-%m-%Y") if p else ""]
            for i, (d, b, g, p) in enumerate(clean)]
    n_clean = len(rows)
    # deliberate mess, inserted back to front so the earlier indexes stay valid; (kind, index to insert before)
    dup = list(rows[55])
    bad = list(rows[70])
    bad[0], bad[3] = "31-09-2026", "SL/26-27/9001"   # no such day
    cn = ["14-08-2026", rows[41][1], "Credit Note", "CN/26-27/0003", inr(-8850), ""]
    plan = [(0, "Grand Total row", None), (70, "Unparseable date", bad), (56, "Duplicate voucher number", dup),
            (41, "Credit note (negative amount)", cn), (23, "Blank row", [""] * len(TALLY_HEAD))]
    for idx, kind, row in plan[1:]:   # the Grand Total row is appended below
        rows.insert(idx, row)
    total = sum(float(r[4].replace(",", "")) for r in rows if r[1] and r[2] == "Sales")
    rows.append(["", "Grand Total", "", "", inr(total), ""])
    # file row numbers: header = row 1, so list index i is row i + 2; recompute from the final list
    where = {"Blank row": next(i for i, r in enumerate(rows) if not any(r)),
             "Credit note (negative amount)": next(i for i, r in enumerate(rows) if r[2] == "Credit Note"),
             "Duplicate voucher number": [i for i, r in enumerate(rows) if r[3] == dup[3]][1],
             "Unparseable date": next(i for i, r in enumerate(rows) if r[0] == "31-09-2026"),
             "Grand Total row": len(rows) - 1}
    notes = [(where[k] + 2, k) for _, k, _ in plan]
    return rows, sorted(notes), n_clean


def write_sample():
    SAMPLES.mkdir(parents=True, exist_ok=True)
    rows, notes, n_clean = tally_rows()
    with open(SAMPLES / "kolhapur_auto_spares_sales_register.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f, lineterminator="\r\n")
        w.writerow(TALLY_HEAD)
        w.writerows(rows)
    wb = Workbook()
    wb.properties.creator, wb.properties.created, wb.properties.modified = "synthetic", pd.Timestamp("2026-10-25").to_pydatetime(), pd.Timestamp("2026-10-25").to_pydatetime()
    ws = wb.active
    ws.title = "Sales Register"
    ws.append(TALLY_HEAD)
    for r in rows:
        ws.append([c if c != "" else None for c in r])   # text cells, as Tally exports them; blank row stays truly empty
    for col, wd in zip("ABCDEF", (12, 24, 12, 16, 16, 12)):
        ws.column_dimensions[col].width = wd
    wb.save(SAMPLES / "kolhapur_auto_spares_sales_register.xlsx")
    ok = [r for r in rows if r[2] == "Sales" and r[0] and r[0] != "31-09-2026"]
    valid = len({r[3] for r in ok})   # distinct vouchers with a good date
    open_ = sum(1 for r in ok if not r[5])
    buyers = sorted({r[1] for r in ok})
    lines = [f"| {n} | {k} |" for n, k in notes]
    (SAMPLES / "README.md").write_text(
        "# Messy upload sample: Kolhapur Auto Spares (fictional, synthetic)\n\n"
        "A TallyPrime-style sales register, same content in `kolhapur_auto_spares_sales_register.xlsx` and `.csv`.\n"
        "Generated by `data/make_demo_tenants.py`; don't edit by hand.\n\n"
        "Columns: `Date` (DD-MM-YYYY text), `Particulars` (party name), `Vch Type`, `Vch No.`, "
        "`Gross Total` (Indian grouping, GST 18% included, e.g. `1,25,000.00`), `Received On` (blank = still open).\n"
        "Row numbers below count the header as row 1 (the same row numbers in Excel and in a text editor).\n\n"
        "## Deliberately bad rows\n\n| Row | What is wrong | \n|---|---|\n" + "\n".join(lines) + "\n\n"
        "The rest should import cleanly. A correct importer finds:\n\n"
        f"- {valid} valid invoices ({open_} open, {valid - open_} paid) after dropping the five rows above "
        f"(the duplicate counts once).\n"
        f"- {len(buyers)} buyers: {', '.join(buyers)}.\n"
        f"- The credit note belongs to {next(r[1] for r in rows if r[2] == 'Credit Note')}; "
        "an importer may subtract it or drop it, but must flag it.\n"
        f"- The file has {len(rows) + 1} rows including the header.\n")
    return n_clean


# ---- self-check ----
def hashes():
    out = {}
    for p in sorted(list(TENANTS.glob("*/*")) + list(SAMPLES.glob("*"))):
        if p.suffix != ".xlsx" and p.name != "owner_entries.json" and "mahesh-cnc" not in p.parts:   # xlsx zips embed write times
            out[str(p.relative_to(ROOT))] = hashlib.sha256(p.read_bytes()).hexdigest()
    return out


def projection(inv, obl, bank, median):
    """13 weekly closing balances: undisputed open invoices land at invoice_date + the buyer's median delay
    (overdue ones a week out), against dated obligations. ponytail: point estimate, no spread; the engine does the 1,000 paths."""
    weeks = []
    open_ = inv[inv.paid_date.isna() & (inv.disputed == 0)]
    land = [max(d + timedelta(days=median[b]), TODAY + timedelta(days=7)) for b, d in zip(open_.buyer, open_.invoice_date)]
    cash = bank
    for k in range(13):
        lo, hi = TODAY + timedelta(days=7 * k + 1), TODAY + timedelta(days=7 * k + 7)
        recv = int(sum(g - t for g, t, d in zip(open_.gross_inr, open_.tds_inr, land) if lo <= d <= hi))
        out = int(obl.amount_inr[(obl.date >= lo) & (obl.date <= hi)].sum())
        cash += recv - out
        weeks.append((lo, recv, out, cash))
    return weeks


def check_tenant(prof, inv, obl, pos, tid):
    d = TENANTS / tid
    assert prof["tenant_id"] == tid and prof["business_name"] and prof["synthetic"] is True and prof["today"] == TODAY.isoformat()
    assert list(pd.read_csv(d / "invoices.csv").columns) == INV_COLS and list(pd.read_csv(d / "obligations.csv").columns) == OBL_COLS
    assert set(pos) == POS_KEYS and all(type(v) is int and v >= 0 for v in pos.values())
    assert set(inv.buyer) <= set(prof["buyers"]) and len(set(prof["buyers"])) == len(prof["buyers"])
    for c in ("invoice_date", "paid_date"):   # valid ISO dates, blank paid_date = open
        for v in inv[c].dropna():
            assert date.fromisoformat(v.isoformat()) == v
    paid = inv.dropna(subset=["paid_date"])
    assert all(p >= i for p, i in zip(paid.paid_date, paid.invoice_date)) and all(p <= TODAY for p in paid.paid_date)
    assert inv.invoice_no.is_unique
    for c in ("taxable_inr", "gst_inr", "gross_inr", "tds_inr", "disputed", "treds_accepted"):
        assert inv[c].dtype.kind == "i", c
    assert (inv.gross_inr == inv.taxable_inr + inv.gst_inr).all()
    assert obl.amount_inr.dtype.kind == "i" and set(obl.kind) <= set(VALID_KINDS) and (obl.amount_inr > 0).all()
    assert obl.date.min() > TODAY and obl.date.max() <= date(2027, 1, 31)
    assert (inv.invoice_date > TODAY).any() and inv[inv.invoice_date > TODAY].paid_date.isna().all()   # planned invoices
    assert (inv.paid_date.isna() & (inv.invoice_date <= TODAY)).any()                                    # open invoices
    return len(paid)


def check_shree(inv, obl, pos):
    assert pos["od_limit_inr"] == 0 and pos["od_drawn_inr"] == 0 and pos["stock_inr"] <= 600_000
    assert set(obl.kind) == {"wages", "supplier", "power", "emi", "rent", "gst", "bonus"}
    assert all(d.day == 20 for d in obl[obl.kind == "gst"].date) and all(d.weekday() == 5 for d in obl[obl.kind == "wages"].date)
    paid = inv.dropna(subset=["paid_date"])
    emp = paid.assign(d=[(p - i).days for p, i in zip(paid.paid_date, paid.invoice_date)]).groupby("buyer").d.median()
    for b in ("Patil Fabrics", "Jain Garments", "Kothari Sarees", "Counter sales"):
        assert abs(emp[b] - SHREE_MEDIAN[b]) <= max(1, 0.2 * SHREE_MEDIAN[b]), (b, emp[b])
    assert emp["Mehta Wholesale"] >= 55   # right-censored (slow payers still open), so the paid-only median reads low
    m = inv.groupby(pd.to_datetime(inv.invoice_date).dt.month).taxable_inr.sum()
    assert min(m[[11, 12, 1, 2]]) > 1.3 * max(m[[6, 7, 8]]), m   # wedding season vs monsoon lull
    weeks = projection(inv, obl, pos["bank_inr"], SHREE_MEDIAN)
    low = min(w[3] for w in weeks)
    print("Shree Textiles: expected receipts (each buyer at median delay) vs obligations, no overdraft; bank Rs %s" % f"{pos['bank_inr']:,}")
    print("  week of      receipts   outgoings     closing")
    for lo, r, o, c in weeks:
        print(f"  {lo}  {r:>9,}  {o:>10,}  {c:>10,}")
    print(f"  lowest week closing: Rs {low:,}")
    assert low < pos["bank_inr"], low   # the season dips below the opening balance (the engine puts the risk at ~23%)
    assert weeks[-1][3] > 0, weeks[-1]       # the season does pull the business back by January
    return weeks


def main():
    write_sample()
    builds = {"shree-textiles": shree(), "sai-snacks": sai()}
    for b in builds.values():
        write_tenant(*b)
    first = hashes()
    for t in builds.values():
        write_tenant(*t)
    write_sample()
    assert hashes() == first, "output differs between runs"   # includes the sample csv and README

    n_paid = {tid: check_tenant(*b, tid) for tid, b in builds.items()}
    check_shree(*builds["shree-textiles"][1:])
    prof, inv, obl, pos = builds["sai-snacks"]
    assert 6 <= n_paid["sai-snacks"] <= 10 and len(prof["buyers"]) in (2, 3), n_paid
    assert (inv.paid_date.isna() & (inv.invoice_date <= TODAY)).sum() >= 3
    # the sample: the xlsx and csv agree, and every deliberate defect is where the README says
    rows, notes, n_clean = tally_rows()
    csv_rows = list(csv.reader(open(SAMPLES / "kolhapur_auto_spares_sales_register.csv", encoding="utf-8")))
    ws = load_workbook(SAMPLES / "kolhapur_auto_spares_sales_register.xlsx").active
    xl_rows = [["" if c is None else str(c) for c in r] for r in ws.iter_rows(values_only=True)]
    assert csv_rows == xl_rows == [TALLY_HEAD] + rows
    assert 60 <= len(rows) <= 120 and len({r[1] for r in rows if r[2] == "Sales"}) == 5
    for n, kind in notes:
        r = csv_rows[n - 1]
        assert {"Blank row": not any(r), "Credit note (negative amount)": r[4].startswith("-"),
                "Duplicate voucher number": r[3] in [x[3] for x in csv_rows[1:n - 1]],
                "Unparseable date": r[0] == "31-09-2026", "Grand Total row": r[1] == "Grand Total"}[kind], (n, kind)
    assert any(r[5] == "" and r[2] == "Sales" and r[0] for r in rows) and "1,25,000.00" == inr(125000)
    print(f"sample: {len(rows)} rows, bad rows at {notes}")
    for tid, (p, i, o, ps) in builds.items():
        print(f"{tid}: {len(i)} invoices ({n_paid[tid]} paid), {len(o)} obligations, bank Rs {ps['bank_inr']:,}, buyers {p['buyers']}")
    print("ok")


if __name__ == "__main__":
    main()
