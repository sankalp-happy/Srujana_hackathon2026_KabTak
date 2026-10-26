"""Onboarding: create a business at sign-up, then fill it from the owner's Kab Tak? Excel workbook.

The workbook has the engine's own layout (docs/planning/multi-tenant-contract.md §1), one sheet per file:
  invoices     invoice_no, buyer, invoice_date, taxable_inr, gst_inr, gross_inr, tds_inr, paid_date, disputed, treds_accepted
  obligations  date, label, kind, amount_inr                                   (optional)
  position     bank_inr, od_limit_inr, od_drawn_inr, stock_inr, creditors_inr  (optional; one row)
Any other sheet (e.g. buyer_summary) is ignored. No LLM: the columns are fixed, so this is plain parsing and checks.
Bad rows are skipped and reported, never guessed.

Check:  python3 ui/ingest.py
"""
import io, json, re, shutil, sys
from datetime import date, datetime
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import pandas as pd
from engine import tenants
from engine.ledger import GST_RATE, KINDS, TDS_RATE

INV_COLS = ["invoice_no", "buyer", "invoice_date", "taxable_inr", "gst_inr", "gross_inr", "tds_inr",
            "paid_date", "disputed", "treds_accepted"]
OBL_COLS = ["date", "label", "kind", "amount_inr"]
POSITION = ["bank_inr", "od_limit_inr", "od_drawn_inr", "stock_inr", "creditors_inr"]
MAX_BUYERS = 20   # the engine copes with about 20; the long tail of small buyers is pooled
OTHER_BUYERS = "Other customers"
MAX_ERRORS = 20   # rows reported back to the page; the rest are only counted


# ---------- cells ----------
def _cell(v):
    if v is None:
        return ""
    if isinstance(v, (datetime, date)):   # Excel date cells
        return (v.date() if isinstance(v, datetime) else v).isoformat()
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _amount(s):
    """'₹1,25,000.00', 'Rs. 4,500', '35600' -> int rupees; None if blank or unreadable."""
    t = re.sub(r"(?i)₹|rs\.?|inr|\s|,", "", str(s))
    try:
        return int(round(float(t))) if t else None
    except ValueError:
        return None


def _date(s):
    """ISO (as in the template) first, then day-first Indian dates (05-04-2026); None if blank or unreadable."""
    s = str(s).strip()
    if not s:
        return None
    try:
        return date.fromisoformat(s[:10])
    except ValueError:
        d = pd.to_datetime(s, dayfirst=True, errors="coerce", format="mixed")
        return None if pd.isna(d) else d.date()


def _nonneg(v, what):
    if isinstance(v, bool) or not isinstance(v, int) or v < 0:
        raise ValueError(f"{what} must be a whole number of rupees, 0 or more.")
    return v


# ---------- sign-up: an empty business ----------
def write_tenant(name, position, owner_name=None, phone=None):
    """Create data/tenants/<new_id>/ per contract §1 with no invoices or bills; returns the tenant id."""
    name = re.sub(r"\s+", " ", str(name or "")).strip()
    if not 0 < len(name) <= 80:
        raise ValueError("Business name is required (80 characters at most).")
    pos = {k: _nonneg(position.get(k, 0), k) for k in POSITION}
    tid = tenants.new_id(name)
    final = tenants.tenant_dir(tid)
    tmp = final.with_name(final.name + ".tmp")
    tmp.mkdir(parents=True)
    try:
        pd.DataFrame(columns=INV_COLS).to_csv(tmp / "invoices.csv", index=False)
        pd.DataFrame(columns=OBL_COLS).to_csv(tmp / "obligations.csv", index=False)
        (tmp / "position.json").write_text(json.dumps(pos, indent=2))
        (tmp / "profile.json").write_text(json.dumps({
            "tenant_id": tid, "business_name": name, "buyers": [], "today": None, "synthetic": False, "params": {},
            **({"owner_name": owner_name} if owner_name else {}), **({"phone": phone} if phone else {})},
            indent=2, ensure_ascii=False))
        tmp.rename(final)
    except BaseException:
        shutil.rmtree(tmp, ignore_errors=True)
        raise
    return tid


# ---------- the Excel workbook ----------
def _sheet(sheets, name, cols):
    """Rows of a sheet as dicts over `cols` (header names matched case-insensitively), with their Excel row number."""
    rows = [[_cell(c) for c in r] for r in sheets[name].iter_rows(values_only=True)]
    start = next((i for i, r in enumerate(rows) if any(rows[i])), None)
    if start is None:
        return []
    head = [h.strip().lower() for h in rows[start]]
    missing = [c for c in cols if c not in head]
    if missing:
        raise ValueError(f"Sheet '{name}' is missing the column{'s' * (len(missing) > 1)} {', '.join(missing)}.")
    at = {c: head.index(c) for c in cols}
    return [(n, {c: (r[at[c]] if at[c] < len(r) else "") for c in cols})
            for n, r in enumerate(rows[start + 1:], start + 2) if any(r)]


def _invoices(rows, errors):
    out, seen = [], set()
    for n, r in rows:
        bad = []
        no, buyer = r["invoice_no"], re.sub(r"\s+", " ", r["buyer"])
        d, paid = _date(r["invoice_date"]), _date(r["paid_date"])
        tax = _amount(r["taxable_inr"])
        if buyer.lower() in ("total", "grand total"):
            continue
        if not no:
            bad.append("missing invoice_no")
        elif no in seen:
            bad.append(f"duplicate invoice_no {no}")
        if not buyer:
            bad.append("missing buyer")
        if d is None:
            bad.append("bad invoice_date")
        if tax is None or tax <= 0:
            bad.append("taxable_inr must be more than 0")
        if r["paid_date"] and paid is None:
            bad.append("bad paid_date")
        elif paid and d and paid < d:
            bad.append("paid before invoiced")
        if bad:
            errors.append({"sheet": "invoices", "row": n, "error": "; ".join(bad)})
            continue
        seen.add(no)
        gst = _amount(r["gst_inr"])
        gst = int(round(tax * GST_RATE)) if gst is None else gst   # blank tax columns: the standard rates
        gross = _amount(r["gross_inr"]) or tax + gst
        tds = _amount(r["tds_inr"])
        out.append({"invoice_no": no, "buyer": buyer, "invoice_date": d.isoformat(), "taxable_inr": tax, "gst_inr": gst,
                    "gross_inr": gross, "tds_inr": int(round(tax * TDS_RATE)) if tds is None else tds,
                    "paid_date": paid.isoformat() if paid else "",
                    "disputed": int(_amount(r["disputed"]) == 1), "treds_accepted": int(_amount(r["treds_accepted"]) == 1)})
    return out


def _bills(rows, errors):
    out = []
    for n, r in rows:
        d, amt, kind = _date(r["date"]), _amount(r["amount_inr"]), r["kind"].strip().lower()
        if d is None or amt is None or amt <= 0:
            errors.append({"sheet": "obligations", "row": n, "error": "needs a date and an amount_inr more than 0"})
            continue
        kind = kind if kind in KINDS else "other"
        out.append({"date": d.isoformat(), "label": r["label"] or kind.replace("_", " ").capitalize(), "kind": kind,
                    "amount_inr": amt})
    return out


def read_workbook(raw_bytes):
    """-> (invoices, obligations, position | None, errors). Raises ValueError if the file can't be used at all."""
    from openpyxl import load_workbook
    try:
        wb = load_workbook(io.BytesIO(raw_bytes), read_only=True, data_only=True)
    except Exception:
        raise ValueError("That isn't an Excel (.xlsx) file we can open.")
    sheets = {ws.title.strip().lower(): ws for ws in wb.worksheets}
    if "invoices" not in sheets:
        raise ValueError("The file needs an 'invoices' sheet, laid out like the Kab Tak? template.")
    errors = []
    inv = _invoices(_sheet(sheets, "invoices", INV_COLS), errors)
    if not inv:
        raise ValueError("No usable invoices in the 'invoices' sheet.")
    obl = _bills(_sheet(sheets, "obligations", OBL_COLS), errors) if "obligations" in sheets else []
    pos = None
    if "position" in sheets and (rows := _sheet(sheets, "position", POSITION)):
        vals = {k: _amount(rows[0][1][k]) for k in POSITION}
        if any(v is None or v < 0 for v in vals.values()):
            errors.append({"sheet": "position", "row": rows[0][0], "error": "needs 5 amounts of 0 or more; kept your sign-up balance"})
        else:
            pos = vals
    return inv, obl, pos, errors


def fill_tenant(tid, inv, obl, pos):
    """Replace the business's invoices and bills with the workbook's (and its position, if the workbook had one).
    Buyers go into the profile by revenue; past the top 19, they are pooled as 'Other customers'."""
    d = tenants.tenant_dir(tid)
    if not tenants.exists(tid):
        raise ValueError(f"Unknown business {tid!r}.")
    df = pd.DataFrame(inv, columns=INV_COLS)
    rev = df.groupby("buyer").taxable_inr.sum().sort_values(ascending=False, kind="stable")
    if len(rev) > MAX_BUYERS:
        df["buyer"] = df.buyer.where(df.buyer.isin(set(rev.index[:MAX_BUYERS - 1])), OTHER_BUYERS)
        rev = df.groupby("buyer").taxable_inr.sum().sort_values(ascending=False, kind="stable")
    df.to_csv(d / "invoices.csv", index=False)
    pd.DataFrame(obl, columns=OBL_COLS).sort_values("date", kind="stable").to_csv(d / "obligations.csv", index=False)
    if pos:
        (d / "position.json").write_text(json.dumps(pos, indent=2))
    prof = json.loads((d / "profile.json").read_text())
    (d / "profile.json").write_text(json.dumps({**prof, "buyers": list(rev.index)}, indent=2, ensure_ascii=False))
    (d / "owner_entries.json").unlink(missing_ok=True)   # the workbook is the new base
    return list(rev.index)


if __name__ == "__main__":
    import tempfile
    from openpyxl import Workbook

    def book(**sheets):
        wb = Workbook()
        wb.remove(wb.active)
        for name, rows in sheets.items():
            ws = wb.create_sheet(name)
            for r in rows:
                ws.append(r)
        b = io.BytesIO()
        wb.save(b)
        return b.getvalue()

    raw = book(
        invoices=[INV_COLS,
                  ["INV-1", "Sankalp", "2024-11-04", 169100, 30438, 199538, 3382, "2025-01-08", 0, 0],
                  ["INV-2", "Rohan", datetime(2026, 10, 20), 18500, None, None, None, None, 0, 1],   # Excel date, blank taxes
                  ["INV-3", "Rohan", "31-02-2026", 5000, 900, 5900, 100, None, 0, 0],                 # bad date
                  ["INV-1", "Rohan", "2026-10-01", 5000, 900, 5900, 100, None, 0, 0],                 # duplicate
                  ["INV-4", "Sankalp", "2026-10-10", 0, 0, 0, 0, None, 0, 0],                         # zero amount
                  ["INV-5", "Aditya", "2026-10-10", "1,20,000", 21600, 141600, 2400, "2026-10-01", 0, 0],   # paid before
                  [None] * 10,
                  ["INV-6", "Aditya", "05-10-2026", "₹20,000", 3600, 23600, 400, "", 1, 0]],         # day-first date
        obligations=[OBL_COLS, ["2026-11-07", "Wages", "wages", 340000], ["2026-11-10", "", "repair", 50000],
                     ["bad", "Rent", "rent", 60000]],
        position=[POSITION, [400000, 4200000, 2380000, 1600000, 998000]],
        buyer_summary=[["buyer", "invoice_count"], ["Sankalp", "=COUNTIF(invoices!B:B,A2)"]])
    inv, obl, pos, errs = read_workbook(raw)
    assert [x["invoice_no"] for x in inv] == ["INV-1", "INV-2", "INV-6"], inv
    assert inv[1] == {"invoice_no": "INV-2", "buyer": "Rohan", "invoice_date": "2026-10-20", "taxable_inr": 18500,
                      "gst_inr": 3330, "gross_inr": 21830, "tds_inr": 370, "paid_date": "", "disputed": 0, "treds_accepted": 1}, inv[1]
    assert inv[2]["invoice_date"] == "2026-10-05" and inv[2]["taxable_inr"] == 20000 and inv[2]["disputed"] == 1
    assert [(e["sheet"], e["row"]) for e in errs] == [("invoices", 4), ("invoices", 5), ("invoices", 6), ("invoices", 7),
                                                      ("obligations", 4)], errs
    assert "paid before invoiced" in errs[3]["error"] and "duplicate" in errs[1]["error"]
    assert obl[1] == {"date": "2026-11-10", "label": "Other", "kind": "other", "amount_inr": 50000}, obl   # unknown kind -> other
    assert pos == {"bank_inr": 400000, "od_limit_inr": 4200000, "od_drawn_inr": 2380000, "stock_inr": 1600000, "creditors_inr": 998000}
    assert read_workbook(book(invoices=[INV_COLS, ["INV-1", "A", "2026-10-01", 100, 18, 118, 2, None, 0, 0]]))[1:3] == ([], None)
    for bad in (b"not excel", book(sales=[["x"]]), book(invoices=[INV_COLS[:-1]]),
                book(invoices=[INV_COLS, ["INV-1", "", "x", 0, 0, 0, 0, None, 0, 0]])):
        try:
            read_workbook(bad)
            raise AssertionError("accepted a bad workbook")
        except ValueError as e:
            print("refused:", e)

    here = Path(__file__).resolve().parent   # the downloadable workbooks (ui/make_template.py) load as they are, no skipped rows
    t_inv, t_obl, t_pos, t_errs = read_workbook((here / "kabtak-template.xlsx").read_bytes())
    assert (len(t_inv), len(t_obl), t_errs) == (2, 2, []) and t_pos and not t_inv[1]["paid_date"] and t_inv[0]["paid_date"]
    mah = here.parent / "data" / "tenants" / "mahesh-cnc"
    m_inv, m_obl, m_pos, m_errs = read_workbook((here / "kabtak-demo.xlsx").read_bytes())
    csv = pd.read_csv(mah / "invoices.csv")
    assert m_errs == [] and len(m_inv) == len(csv) and sum(x["taxable_inr"] for x in m_inv) == csv.taxable_inr.sum()
    assert sum(not x["paid_date"] for x in m_inv) == csv.paid_date.isna().sum() and len(m_obl) == len(pd.read_csv(mah / "obligations.csv"))
    assert m_pos == json.loads((mah / "position.json").read_text()) and {x["buyer"] for x in m_inv} == set(json.loads((mah / "profile.json").read_text())["buyers"])

    tenants.DATA_ROOT = Path(tempfile.mkdtemp())   # same effect as KABTAK_DATA=<tmp>: tenant_dir reads it per call
    tid = write_tenant("Patil Steel Works", {"bank_inr": 250000}, owner_name="Ramesh", phone="919876543210")
    d = tenants.tenant_dir(tid)
    assert list(pd.read_csv(d / "invoices.csv").columns) == INV_COLS and json.loads((d / "position.json").read_text())["bank_inr"] == 250000
    assert fill_tenant(tid, inv, obl, None) == ["Sankalp", "Aditya", "Rohan"]   # by revenue
    assert json.loads((d / "position.json").read_text())["bank_inr"] == 250000   # no position sheet: sign-up balance stays
    fill_tenant(tid, inv, obl, pos)
    assert json.loads((d / "position.json").read_text()) == pos and list(pd.read_csv(d / "invoices.csv").columns) == INV_COLS
    prof = json.loads((d / "profile.json").read_text())
    assert prof["buyers"] == ["Sankalp", "Aditya", "Rohan"] and prof["owner_name"] == "Ramesh" and prof["today"] is None
    many = [{**inv[0], "buyer": f"B{i}", "taxable_inr": 1000 + i, "invoice_no": f"X{i}"} for i in range(30)]
    assert len(fill_tenant(tid, many, [], None)) == MAX_BUYERS and OTHER_BUYERS in fill_tenant(tid, many, [], None)
    for bad in (lambda: write_tenant("", {}), lambda: write_tenant("x", {"bank_inr": -1}), lambda: fill_tenant("nope-0000", inv, [], None)):
        try:
            bad()
            raise AssertionError("accepted bad input")
        except ValueError:
            pass
    print("ok")
