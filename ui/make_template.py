"""Writes the two downloadable workbooks from ingest's column constants: kabtak-template.xlsx (empty, with example
rows) and kabtak-demo.xlsx (the mahesh-cnc demo ledger, so a demo always has a file that uploads cleanly).

Run from the repo root:  python3 ui/make_template.py     (ui/ingest.py's self-check asserts both still load)
"""
import json
import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import pandas as pd
from openpyxl import Workbook
from ingest import INV_COLS, OBL_COLS, POSITION

UI = Path(__file__).resolve().parent
MAHESH = UI.parent / "data" / "tenants" / "mahesh-cnc"
HOW = ["Fill the invoices, obligations and position sheets, then upload this file on the Kab Tak? sign-up page.",
       "invoices: one row per invoice you raised. Leave paid_date empty while it is unpaid. dates as 2026-10-20 or 20-10-2026.",
       "obligations: money you must pay out (kind: wages, emi, gst, rent, power, advance_tax, bonus, supplier, other).",
       "position: one row, in rupees. Delete the example rows first. Other sheets are ignored."]


def book(path, invoices, obligations, position):
    wb = Workbook()
    wb.remove(wb.active)
    for name, cols, rows in (("invoices", INV_COLS, invoices), ("obligations", OBL_COLS, obligations), ("position", POSITION, [position])):
        ws = wb.create_sheet(name)
        ws.append(cols)
        for r in rows:
            ws.append(r)
    wb.create_sheet("how_to").append([HOW[0]])
    for line in HOW[1:]:
        wb["how_to"].append([line])
    wb.save(path)


def csv_rows(f, cols):   # blank cells stay empty, numbers stay numbers
    return [[None if pd.isna(v) else v.item() if hasattr(v, "item") else v for v in r] for r in pd.read_csv(f)[cols].itertuples(index=False)]


if __name__ == "__main__":
    book(UI / "kabtak-template.xlsx",
         [["INV-001", "Example Buyer", "2026-09-01", 100000, 18000, 118000, 2000, "2026-10-10", 0, 0],
          ["INV-002", "Example Buyer", "2026-10-15", 50000, 9000, 59000, 1000, None, 0, 0]],
         [["2026-11-07", "Wages", "wages", 150000], ["2026-11-20", "GST", "gst", 40000]],
         [100000, 500000, 300000, 200000, 80000])
    pos = json.loads((MAHESH / "position.json").read_text())
    book(UI / "kabtak-demo.xlsx", csv_rows(MAHESH / "invoices.csv", INV_COLS), csv_rows(MAHESH / "obligations.csv", OBL_COLS),
         [pos[k] for k in POSITION])
    print("wrote ui/kabtak-template.xlsx, ui/kabtak-demo.xlsx")
